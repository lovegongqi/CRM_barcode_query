"""Price-free Excel reports for inventory stocktake differences."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


SUMMARY_HEADERS = [
    "商品条码", "商品名称", "规格", "型号", "类别", "单位", "序列号管理",
    "完成时账面数量", "实盘数量", "盘盈/盘亏", "处理状态", "商品备注",
]
SERIAL_HEADERS = [
    "商品条码", "商品名称", "序列号", "差异分类", "盘点任务", "扫描账号",
    "扫描设备", "扫描时间", "处理状态", "序列号备注",
]

_SUMMARY_WIDTHS = (18, 24, 14, 14, 14, 10, 12, 18, 14, 14, 14, 30)
_SERIAL_WIDTHS = (18, 24, 24, 20, 22, 14, 16, 22, 14, 30)

_KIND_LABELS = {
    "product_quantity": "商品数量差异",
    "system_only_serial": "账面独有",
    "physical_only_serial": "实物独有",
    "other_product_serial": "其他商品",
    "already_shipped_serial": "已出库序列号",
    "unknown_serial": "未知序列号",
    "serial_unverified": "序列号未核对",
}


def _text(value):
    """Return a string; workbook finalization marks it as a literal cell."""
    if value is None:
        return ""
    value = str(value)
    return value


def _value(row, *names, default=""):
    for name in names:
        if name in row and row[name] is not None:
            return row[name]
    return default


def _style_sheet(sheet, headers, widths, text_columns=()):
    sheet.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{chr(64 + len(headers))}1"
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for column in text_columns:
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for value_cell in cell:
                value_cell.number_format = "@"


def _save(workbook):
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    # Explicitly force all user-controlled strings (including
                    # values beginning with =,+,-,@) to remain text cells.
                    cell.data_type = "s"
                    cell.number_format = "@"
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _summary_row(item, discrepancy):
    discrepancy_kind = _value(discrepancy or {}, "kind")
    return [
        _text(_value(item, "barcode")),
        _text(_value(item, "name")),
        _text(_value(item, "spec")),
        _text(_value(item, "model")),
        _text(_value(item, "category")),
        _text(_value(item, "unit")),
        _text("是" if _value(item, "has_serial", default=False) else "否"),
        _text(_value(item, "completed_book_qty", "completed_book_quantity", "book_quantity")),
        _text(_value(item, "completed_actual_qty", "completed_counted_quantity", "actual_quantity")),
        _text(_value(item, "diff_qty", "difference", "difference_qty")),
        _text(
            _KIND_LABELS.get(discrepancy_kind)
            if discrepancy_kind == "serial_unverified"
            else _value(
                discrepancy or {}, "state",
                default=_value(item, "state", "status"),
            )
        ),
        _text(_value(discrepancy or {}, "note", default=_value(item, "note", "remark"))),
    ]


def _serial_row(row, state_label=None, task_id=""):
    return [
        _text(_value(row, "barcode")),
        _text(_value(row, "name", "lookup_name")),
        _text(_value(row, "serial")),
        _text(_KIND_LABELS.get(
            _value(row, "kind", "classification"),
            _value(row, "kind", "classification"),
        )),
        _text(_value(row, "task_id", default=task_id)),
        _text(_value(row, "scan_actor", "actor")),
        _text(_value(row, "scan_device", "device_id")),
        _text(_value(row, "scanned_at", "scan_time")),
        _text(state_label if state_label is not None else _value(row, "state", "status")),
        _text(_value(row, "note", "serial_note", "remark")),
    ]


def build_inventory_workbook(task, items, discrepancies):
    """Build the two-sheet task report using only explicitly allowlisted fields."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "商品差异汇总"
    serials = workbook.create_sheet("序列号差异明细")
    _style_sheet(summary, SUMMARY_HEADERS, _SUMMARY_WIDTHS, (1,))
    _style_sheet(serials, SERIAL_HEADERS, _SERIAL_WIDTHS, (1, 3))

    product_discrepancies = {}
    for discrepancy in discrepancies or []:
        barcode = _value(discrepancy, "barcode")
        if _value(discrepancy, "serial") in ("", None) and barcode:
            if (
                _value(discrepancy, "kind") == "serial_unverified"
                or barcode not in product_discrepancies
            ):
                product_discrepancies[barcode] = discrepancy
    for item in items or []:
        difference = _value(item, "diff_qty", "difference", "difference_qty", default="0")
        if str(difference) == "0":
            continue
        summary.append(_summary_row(item, product_discrepancies.get(_value(item, "barcode"))))

    for row in discrepancies or []:
        if (
            _value(row, "serial") in ("", None)
            and _value(row, "kind") != "serial_unverified"
        ):
            continue
        serials.append(_serial_row(row, task_id=_value(task or {}, "task_id")))
    return _save(workbook)


def build_discrepancy_workbook(rows, state_label):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "差异明细"
    _style_sheet(sheet, SERIAL_HEADERS, _SERIAL_WIDTHS, (1, 3))
    for row in rows or []:
        sheet.append(_serial_row(row, state_label=state_label))
    return _save(workbook)
