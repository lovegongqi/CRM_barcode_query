import atexit
from datetime import datetime, timedelta
from io import BytesIO
import os
import tempfile
import threading
import unittest
from unittest import mock

from openpyxl import load_workbook


TEST_DATA_DIR = tempfile.TemporaryDirectory()
atexit.register(TEST_DATA_DIR.cleanup)
os.environ["CRM_DATA_DIR"] = TEST_DATA_DIR.name
os.environ["CRM_DESKTOP_APP"] = "0"

import app as app_module
from inventory_service import InventoryService, InventoryServiceError
from inventory_store import (
    InventoryConfirmationRequired,
    InventoryConflict,
    InventoryNotFound,
    InventoryPermissionDenied,
    InventoryStore,
    InventoryVersionConflict,
)


class InventoryRouteTest(unittest.TestCase):
    def setUp(self):
        app_module.app.config.update(TESTING=True)
        self.tempdir = tempfile.TemporaryDirectory()
        self.accounts_file = os.path.join(self.tempdir.name, "accounts.json")
        self.accounts_patch = mock.patch.object(
            app_module, "ACCOUNTS_FILE", self.accounts_file
        )
        self.accounts_patch.start()
        app_module.save_accounts([
            {
                "id": "admin",
                "username": "admin",
                "display_name": "管理员",
                "password": "admin-pass",
                "permissions": [
                    "crm", "results", "transfer", "inbound", "inventory",
                    "accounts", "product-library",
                ],
                "updated_at": "",
            },
            {
                "id": "viewer-id",
                "username": "viewer",
                "display_name": "访客",
                "password": "viewer-pass",
                "permissions": ["results"],
                "updated_at": "",
            },
            {
                "id": "counter-id",
                "username": "counter",
                "display_name": "盘点员",
                "password": "counter-pass",
                "permissions": ["inventory"],
                "updated_at": "",
            },
            {
                "id": "warehouse-account-id",
                "username": "warehouse-user",
                "display_name": "仓库账号",
                "password": "warehouse-pass",
                "permissions": ["inbound", "inventory"],
                "updated_at": "",
            },
        ])
        self.store = mock.Mock()
        self.service = mock.Mock()
        self.production_inventory_service = app_module.inventory_service
        self.store.get_active_task.return_value = None
        self.store.get_task_version.return_value = 3
        self.store.get_task_snapshot.return_value = {
            "success": True,
            "task_id": "task-1",
            "phase": "counting",
            "version": 3,
            "last_sync_at": datetime.now().isoformat(timespec="seconds"),
            "items": [],
        }
        self.store.list_items.return_value = []
        self.store.list_task_history.return_value = {
            "tasks": [], "total": 0, "limit": 20, "offset": 0,
        }
        self.store.list_discrepancies.return_value = []
        self.store.list_audit_events.return_value = [{
            "id": 7, "barcode": "A/B", "event_type": "count_entry_updated",
            "event_label": "修改分次数量", "actor": "counter",
            "device_id": "device-a", "created_at": "2026-09-05T10:00:00",
            "before_quantity": "12", "after_quantity": "13",
        }]
        self.service.create_task.return_value = {"task_id": "task-1", "phase": "counting"}
        self.service.open_count_item.return_value = {"barcode": "A/B", "state": "pending"}
        self.store.heartbeat_lock.return_value = {"barcode": "A/B", "device_id": "device-a"}
        self.service.submit_count.return_value = {"barcode": "A/B", "state": "matched"}
        self.service.add_count_entry.return_value = {
            "barcode": "A/B", "count_total": "12", "task_version": 4,
        }
        self.service.update_count_entry.return_value = {
            "barcode": "A/B", "count_total": "13", "task_version": 5,
        }
        self.service.delete_count_entry.return_value = {
            "barcode": "A/B", "count_total": None, "task_version": 6,
        }
        self.service.open_serial_item.return_value = {"barcode": "A/B", "counts": {}}
        self.service.refresh_serial_item.return_value = {"barcode": "A/B", "counts": {}}
        self.service.scan_serial.return_value = {"serial": "SN/1", "classification": "matched"}
        self.service.delete_serial_scan.return_value = {"barcode": "A/B", "removed": "SN/1"}
        self.service.finish_serial_item.return_value = {"barcode": "A/B", "counts": {}}
        self.store.get_carton_preset.return_value = {
            "barcode": "A/B", "carton_quantity": 20,
        }
        self.service.save_carton_preset.return_value = {
            "barcode": "A/B", "carton_quantity": 20, "version": 4,
        }
        carton_serial = {
            "barcode": "A/B", "cartons": [], "counts": {}, "version": 4,
        }
        self.service.create_serial_carton.return_value = carton_serial
        self.service.add_carton_serial.return_value = carton_serial
        self.service.remove_carton_serial.return_value = carton_serial
        self.service.delete_serial_carton.return_value = carton_serial
        self.service.complete_task.return_value = {"task_id": "task-1", "completed": True}
        self.service.reopen_task.return_value = {
            "task_id": "task-1", "phase": "counting", "version": 9,
        }
        self.store.admin_unlock.return_value = {"task_id": "task-1", "barcode": "A/B", "unlocked": True}
        self.store.add_discrepancy_note.return_value = {
            "id": 1, "note": "checked", "task_version": 3,
        }
        self.store.archive_discrepancy.return_value = {
            "id": 1, "state": "archived", "task_version": 3,
        }
        self.store.restore_discrepancy.return_value = {
            "id": 1, "state": "open", "task_version": 3,
        }
        self.store_patch = mock.patch.object(
            app_module, "inventory_store", self.store, create=True
        )
        self.service_patch = mock.patch.object(
            app_module, "inventory_service", self.service, create=True
        )
        self.store_patch.start()
        self.service_patch.start()
        if hasattr(app_module, "inventory_sync_threads"):
            app_module.inventory_sync_threads.clear()
        if hasattr(app_module, "inventory_serial_prefetch_threads"):
            app_module.inventory_serial_prefetch_threads.clear()

    def tearDown(self):
        if hasattr(app_module, "inventory_sync_threads"):
            for thread in list(app_module.inventory_sync_threads.values()):
                thread.join(timeout=1)
            app_module.inventory_sync_threads.clear()
        if hasattr(app_module, "inventory_serial_prefetch_threads"):
            for thread in list(
                app_module.inventory_serial_prefetch_threads.values()
            ):
                thread.join(timeout=1)
            app_module.inventory_serial_prefetch_threads.clear()
        self.service_patch.stop()
        self.store_patch.stop()
        self.accounts_patch.stop()
        self.tempdir.cleanup()

    def login_account(self, username, password=None):
        client = app_module.app.test_client()
        response = client.post(
            "/api/app-auth/login",
            json={"username": username, "password": password or f"{username}-pass"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["success"])
        return client

    def test_inventory_permission_controls_page_and_api(self):
        viewer = self.login_account("viewer")
        self.assertEqual(viewer.get("/inventory").status_code, 403)
        self.assertEqual(viewer.get("/api/inventory/tasks/active").status_code, 403)

        counter = self.login_account("counter")
        self.assertEqual(counter.get("/inventory").status_code, 200)
        response = counter.get("/api/inventory/tasks/active")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json(), {"success": True, "task": None})

    def test_anonymous_inventory_page_and_api_require_login(self):
        client = app_module.app.test_client()
        page = client.get("/inventory", follow_redirects=False)
        api = client.get("/api/inventory/tasks/active")
        self.assertEqual(page.status_code, 302)
        self.assertTrue(page.headers["Location"].startswith("/login?next="))
        self.assertEqual(api.status_code, 401)

    def test_admin_navigation_places_inventory_after_inbound(self):
        page = self.login_account("admin", "admin-pass").get("/inventory").get_data(as_text=True)
        self.assertLess(page.index('href="/inbound"'), page.index('href="/inventory"'))
        self.assertLess(page.index('href="/inventory"'), page.index('href="/product-library"'))

    def test_inventory_page_receives_only_public_current_account(self):
        page = self.login_account("counter").get("/inventory").get_data(as_text=True)
        self.assertIn('data-aurora-page="inventory"', page)
        self.assertIn("const CURRENT_ACCOUNT =", page)
        self.assertIn('"is_admin": false', page)
        self.assertNotIn("counter-pass", page)
        self.assertNotIn('value="archive"', page)

    def test_inventory_page_uses_filter_instead_of_duplicate_serial_queue(self):
        page = self.login_account("counter").get("/inventory").get_data(as_text=True)
        self.assertIn('<option value="serial_pending">待序列号</option>', page)
        self.assertNotIn('id="inventorySerialQueueRoot"', page)
        self.assertNotIn('id="inventorySerialQueue"', page)

    def test_admin_default_and_account_save_allow_inventory_permission(self):
        os.remove(self.accounts_file)
        self.assertIn("inventory", app_module.load_accounts()[0]["permissions"])

        client = self.login_account("admin", "88293529")
        response = client.post(
            "/api/accounts",
            json={
                "username": "new-counter",
                "display_name": "新盘点员",
                "password": "secret",
                "permissions": ["inventory", "not-allowed"],
            },
        )
        self.assertEqual(response.status_code, 200)
        saved = next(row for row in app_module.load_accounts() if row["username"] == "new-counter")
        self.assertEqual(saved["permissions"], ["inventory"])

    def test_inventory_service_uses_shared_gyj_business_provider_for_account_alias(self):
        pool = mock.Mock()
        with mock.patch.object(app_module, "gyj_worker", pool):
            result = self.production_inventory_service.worker_provider("warehouse-account-id")
        self.assertIsInstance(result, app_module.GYJBusinessWorker)
        self.assertIs(result.pool, pool)
        self.assertEqual(result.actor, "warehouse-user")
        pool.get.assert_not_called()

    def test_create_uses_shared_owner_and_session_actor_not_request_overrides(self):
        client = self.login_account("warehouse-user", "warehouse-pass")
        response = client.post(
            "/api/inventory/tasks",
            json={"owner": "admin", "actor": "spoofed", "device_id": "ignored"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(set(response.get_json()), {"success", "task", "version"})
        self.service.create_task.assert_called_once_with(
            "admin", "warehouse-user"
        )

    def test_active_stocktake_is_shared_across_inventory_accounts(self):
        self.store.get_active_task.return_value = {"task_id": "task-1"}
        with mock.patch.object(app_module, "_ensure_inventory_sync"):
            admin_response = self.login_account("admin", "admin-pass").get(
                "/api/inventory/tasks/active"
            )
            counter_response = self.login_account("counter").get(
                "/api/inventory/tasks/active"
            )

        self.assertEqual(admin_response.status_code, 200)
        self.assertEqual(counter_response.status_code, 200)
        self.assertEqual(
            self.store.get_active_task.call_args_list,
            [mock.call("admin"), mock.call("admin")],
        )
        self.assertEqual(
            [call.args[:2] for call in self.store.get_task_snapshot.call_args_list],
            [("admin", "task-1"), ("admin", "task-1")],
        )

    def test_real_stocktake_and_history_are_shared_but_actor_is_preserved(self):
        class SharedWorker:
            def load_inventory_catalog(self):
                return True, [{
                    "barcode": "A1", "name": "Filter", "spec": "",
                    "model": "", "category": "Parts", "unit": "piece",
                    "has_serial": False, "initial_stock": "2",
                }]

            def read_inventory_stock(self, _barcode):
                return True, "2"

        store = InventoryStore(
            os.path.join(self.tempdir.name, "shared-route.sqlite3")
        )
        service = InventoryService(store, lambda _owner: SharedWorker())
        with (
            mock.patch.object(app_module, "inventory_store", store),
            mock.patch.object(app_module, "inventory_service", service),
            mock.patch.object(app_module, "_ensure_inventory_sync"),
        ):
            admin = self.login_account("admin", "admin-pass")
            created = admin.post("/api/inventory/tasks").get_json()
            task_id = created["task"]["task_id"]

            counter = self.login_account("counter")
            active = counter.get("/api/inventory/tasks/active").get_json()
            self.assertEqual(active["task"]["task_id"], task_id)
            counted = counter.post(
                f"/api/inventory/tasks/{task_id}/items/A1/count-entries",
                json={"device_id": "device-counter", "quantity": "2"},
            ).get_json()

            completed = admin.post(
                f"/api/inventory/tasks/{task_id}/complete",
                json={
                    "expected_version": counted["version"],
                    "allow_unverified_serials": False,
                },
            )
            self.assertEqual(completed.status_code, 200)
            history = counter.get("/api/inventory/tasks/history").get_json()

        self.assertEqual(history["tasks"][0]["task_id"], task_id)
        events = store.list_audit_events("admin", task_id, barcode="A1")
        self.assertIn("counter", {event["actor"] for event in events})

    def test_active_query_passes_version_query_and_state_to_store(self):
        self.store.get_active_task.return_value = {"task_id": "task-1"}
        client = self.login_account("counter")
        with mock.patch.object(app_module, "_ensure_inventory_sync") as ensure:
            response = client.get(
                "/api/inventory/tasks/active?version=7&query=Zero%20stock&state=pending"
            )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(set(payload), {"success", "task"})
        self.store.get_task_snapshot.assert_called_once_with(
            "admin", "task-1", known_version=7
        )
        self.store.list_items.assert_called_once_with(
            "task-1", query="Zero stock", state="pending", include_zero=True
        )
        ensure.assert_called_once_with("admin", "task-1")

    def test_active_unchanged_snapshot_does_not_reload_items(self):
        self.store.get_active_task.return_value = {"task_id": "task-1"}
        self.store.get_task_snapshot.return_value = {
            "success": True, "unchanged": True, "version": 7
        }
        with mock.patch.object(app_module, "_ensure_inventory_sync"):
            response = self.login_account("counter").get(
                "/api/inventory/tasks/active?version=7"
            )
        self.assertEqual(
            response.get_json()["task"],
            {"success": True, "unchanged": True, "version": 7},
        )
        self.store.list_items.assert_not_called()

    def test_active_task_returns_unfiltered_categories(self):
        self.store.get_active_task.return_value = {"task_id": "task-1"}
        self.store.get_task_snapshot.return_value = {
            "task_id": "task-1", "phase": "counting", "version": 7,
            "last_sync_at": datetime.now().isoformat(timespec="seconds"),
            "items": [
                {"barcode": "A", "category": "滤芯"},
                {"barcode": "B", "category": "整机"},
                {"barcode": "C", "category": "滤芯"},
                {"barcode": "D", "category": ""},
            ],
        }
        self.store.list_items.return_value = [{
            "barcode": "B", "category": "整机",
        }]

        with mock.patch.object(app_module, "_ensure_inventory_sync"):
            response = self.login_account("counter").get(
                "/api/inventory/tasks/active?query=B"
            )

        task = response.get_json()["task"]
        self.assertEqual(task["categories"], ["整机", "滤芯"])
        self.assertEqual([item["barcode"] for item in task["items"]], ["B"])

    def test_all_json_routes_use_the_declared_methods_and_response_keys(self):
        client = self.login_account("admin", "admin-pass")
        cases = [
            ("post", "/api/inventory/tasks", {}, "task"),
            ("get", "/api/inventory/tasks/history", None, "tasks"),
            ("get", "/api/inventory/tasks/task-1", None, "task"),
            ("post", "/api/inventory/tasks/task-1/items/A%2FB/serial/open", {"device_id": "device-a"}, "serial"),
            ("post", "/api/inventory/tasks/task-1/items/A%2FB/serial/refresh", {"device_id": "device-a"}, "serial"),
            ("post", "/api/inventory/tasks/task-1/items/A%2FB/serials", {"device_id": "device-a", "serial": "SN/1"}, "scan"),
            ("delete", "/api/inventory/tasks/task-1/items/A%2FB/serials/SN%2F1", {"device_id": "device-a"}, "serial"),
            ("post", "/api/inventory/tasks/task-1/items/A%2FB/serial/finish", {"device_id": "device-a"}, "serial"),
            ("post", "/api/inventory/tasks/task-1/complete", {"expected_version": 3, "allow_unverified_serials": False}, "task"),
            ("post", "/api/inventory/tasks/task-1/items/A%2FB/unlock", {"is_admin": False, "expected_version": 3}, "result"),
            ("get", "/api/inventory/discrepancies?state=open&query=A%2FB", None, "discrepancies"),
            ("post", "/api/inventory/discrepancies/1/notes", {"note": "checked", "serial": "SN/1", "actor": "spoofed", "expected_version": 3}, "note"),
            ("post", "/api/inventory/discrepancies/1/archive", {"is_admin": False, "expected_version": 3}, "discrepancy"),
            ("post", "/api/inventory/discrepancies/1/restore", {"is_admin": False, "expected_version": 3}, "discrepancy"),
        ]
        for method, path, body, key in cases:
            with self.subTest(method=method, path=path):
                response = getattr(client, method)(path, json=body)
                self.assertEqual(response.status_code, 200)
                expected_keys = {"success", key}
                if path == "/api/inventory/tasks/history":
                    expected_keys.add("pagination")
                if method in {"post", "delete"}:
                    expected_keys.add("version")
                self.assertEqual(set(response.get_json()), expected_keys)
                self.assertTrue(response.get_json()["success"])

        self.service.open_serial_item.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "admin",
        )
        self.service.refresh_serial_item.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "admin", force=False,
        )
        self.service.scan_serial.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "admin", "SN/1",
        )
        self.service.delete_serial_scan.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "admin", "SN/1",
        )
        self.service.finish_serial_item.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "admin",
        )
        self.service.complete_task.assert_called_once_with(
            "admin", "task-1", "admin",
            allow_unverified_serials=False, expected_version=3,
        )
        self.store.admin_unlock.assert_called_once_with(
            "task-1", "A/B", "admin", True, expected_version=3,
        )
        self.store.add_discrepancy_note.assert_called_once_with(
            "admin", 1, "admin", "checked", serial="SN/1",
            expected_version=3,
        )
        self.store.archive_discrepancy.assert_called_once_with(
            "admin", 1, "admin", True, expected_version=3,
        )
        self.store.restore_discrepancy.assert_called_once_with(
            "admin", 1, "admin", True, expected_version=3,
        )

    def test_partial_count_entry_routes(self):
        client = self.login_account("counter")
        self.service.open_count_item.return_value = {
            "barcode": "A/B", "count_entries": [], "count_total": None,
            "task_version": 3,
        }

        opened = client.get(
            "/api/inventory/tasks/task-1/items/A%2FB/count-entries"
        )
        added = client.post(
            "/api/inventory/tasks/task-1/items/A%2FB/count-entries",
            json={"device_id": "device-a", "quantity": "12"},
        )
        updated = client.post(
            "/api/inventory/tasks/task-1/items/A%2FB/count-entries/7",
            json={
                "device_id": "device-b", "quantity": "13",
                "entry_version": 1,
            },
        )
        deleted = client.delete(
            "/api/inventory/tasks/task-1/items/A%2FB/count-entries/7",
            json={"device_id": "device-a", "entry_version": 2},
        )

        self.assertEqual(opened.status_code, 200)
        self.assertEqual(opened.get_json()["item"]["barcode"], "A/B")
        self.assertEqual(added.get_json()["version"], 4)
        self.assertEqual(updated.get_json()["version"], 5)
        self.assertEqual(deleted.get_json()["version"], 6)
        self.service.open_count_item.assert_called_once_with(
            "admin", "task-1", "A/B", "counter"
        )
        self.service.add_count_entry.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "counter", "12"
        )
        self.service.update_count_entry.assert_called_once_with(
            "admin", "task-1", "A/B", 7, 1,
            "device-b", "counter", "13",
        )
        self.service.delete_count_entry.assert_called_once_with(
            "admin", "task-1", "A/B", 7, 2,
            "device-a", "counter",
        )

        for path in (
            "/api/inventory/tasks/task-1/items/A%2FB/claim",
            "/api/inventory/tasks/task-1/items/A%2FB/heartbeat",
            "/api/inventory/tasks/task-1/items/A%2FB/count",
        ):
            response = client.post(path, json={})
            self.assertEqual(response.status_code, 409)
            self.assertTrue(response.get_json()["refresh_required"])

    def test_serialized_count_save_schedules_background_serial_prefetch(self):
        self.service.add_count_entry.side_effect = [
            {
                "barcode": "B2", "has_serial": True,
                "state": "serial_pending", "count_total": "1", "task_version": 4,
            },
            {
                "barcode": "A1", "has_serial": False,
                "state": "variance", "count_total": "1", "task_version": 5,
            },
        ]
        client = self.login_account("counter")
        with mock.patch.object(
            app_module, "_ensure_inventory_serial_prefetch", create=True
        ) as ensure_prefetch:
            serial_response = client.post(
                "/api/inventory/tasks/task-1/items/B2/count-entries",
                json={"device_id": "device-a", "quantity": "1"},
            )
            ordinary_response = client.post(
                "/api/inventory/tasks/task-1/items/A1/count-entries",
                json={"device_id": "device-a", "quantity": "1"},
            )

        self.assertEqual(serial_response.status_code, 200)
        self.assertEqual(ordinary_response.status_code, 200)
        ensure_prefetch.assert_called_once_with(
            "admin", "task-1", "B2", "device-a"
        )

    def test_serialized_count_delete_schedules_background_serial_prefetch(self):
        self.service.delete_count_entry.return_value = {
            "barcode": "B2", "has_serial": True,
            "state": "serial_pending", "count_total": "1", "task_version": 6,
        }
        client = self.login_account("counter")
        with mock.patch.object(
            app_module, "_ensure_inventory_serial_prefetch", create=True
        ) as ensure_prefetch:
            response = client.delete(
                "/api/inventory/tasks/task-1/items/B2/count-entries/7",
                json={
                    "device_id": "device-a", "entry_version": 2,
                },
            )

        self.assertEqual(response.status_code, 200)
        ensure_prefetch.assert_called_once_with(
            "admin", "task-1", "B2", "device-a"
        )

    def test_retired_quantity_lock_routes_request_page_refresh(self):
        client = self.login_account("counter")
        for suffix in ("claim", "heartbeat", "count"):
            response = client.post(
                f"/api/inventory/tasks/task-1/items/A/{suffix}", json={}
            )
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.get_json()["refresh_required"], True)
        self.service.submit_count.assert_not_called()
        self.store.heartbeat_lock.assert_not_called()

    def test_version_conflict_returns_current_version_without_changing_plain_conflicts(self):
        client = self.login_account("counter")
        self.service.update_count_entry.side_effect = InventoryVersionConflict(
            "盘点任务已被其他设备更新", current_version=9
        )
        stale = client.post(
            "/api/inventory/tasks/task-1/items/A/count-entries/7",
            json={
                "device_id": "device-a", "quantity": "2",
                "entry_version": 1,
            },
        )
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(stale.get_json(), {
            "success": False,
            "error": "盘点任务已被其他设备更新",
            "current_version": 9,
        })

        self.service.update_count_entry.side_effect = InventoryConflict("当前阶段不允许此操作")
        ordinary = client.post(
            "/api/inventory/tasks/task-1/items/A/count-entries/7",
            json={
                "device_id": "device-a", "quantity": "2",
                "entry_version": 1,
            },
        )
        self.assertEqual(ordinary.status_code, 409)
        self.assertNotIn("current_version", ordinary.get_json())

    def test_completion_confirmation_is_structured_and_boolean_is_validated(self):
        client = self.login_account("counter")
        self.service.complete_task.side_effect = InventoryConfirmationRequired(
            "仍有 2 个商品未核对序列号",
            pending_serial_count=2,
        )

        warning = client.post(
            "/api/inventory/tasks/task-1/complete",
            json={"expected_version": 3, "allow_unverified_serials": False},
        )
        self.assertEqual(warning.status_code, 409)
        self.assertEqual(warning.get_json(), {
            "success": False,
            "confirmation_required": True,
            "pending_serial_count": 2,
            "error": "仍有 2 个商品未核对序列号",
        })

        self.service.complete_task.reset_mock(side_effect=True)
        invalid = client.post(
            "/api/inventory/tasks/task-1/complete",
            json={"allow_unverified_serials": "true"},
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("allow_unverified_serials", invalid.get_json()["error"])
        self.service.complete_task.assert_not_called()

    def test_stale_key_mutations_return_409_and_retry_with_refreshed_version(self):
        client = self.login_account("admin", "admin-pass")
        cases = [
            (
                "post", "/api/inventory/tasks/task-1/items/A/count-entries/7",
                {"device_id": "device-b", "quantity": "2", "entry_version": 1},
                self.service.update_count_entry, "item",
                {"barcode": "A", "state": "matched", "task_version": 10},
            ),
            (
                "post", "/api/inventory/tasks/task-1/items/A/serials",
                {"device_id": "device-b", "serial": "SN-1"},
                self.service.scan_serial, "scan",
                {"serial": "SN-1", "classification": "matched", "task_version": 10},
            ),
            (
                "delete", "/api/inventory/tasks/task-1/items/A/serials/SN-1",
                {"device_id": "device-b"}, self.service.delete_serial_scan,
                "serial", {"barcode": "A", "removed": "SN-1", "task_version": 10},
            ),
            (
                "post", "/api/inventory/discrepancies/1/notes",
                {"note": "复核"}, self.store.add_discrepancy_note,
                "note", {"id": 1, "note": "复核", "task_version": 10},
            ),
            (
                "post", "/api/inventory/discrepancies/1/archive", {},
                self.store.archive_discrepancy, "discrepancy",
                {"id": 1, "state": "archived", "task_version": 10},
            ),
            (
                "post", "/api/inventory/discrepancies/1/restore", {},
                self.store.restore_discrepancy, "discrepancy",
                {"id": 1, "state": "open", "task_version": 10},
            ),
            (
                "post", "/api/inventory/tasks/task-1/items/A/unlock", {},
                self.store.admin_unlock, "result",
                {"task_id": "task-1", "barcode": "A", "unlocked": True, "task_version": 10},
            ),
        ]
        for method, path, body, target, key, success_value in cases:
            with self.subTest(path=path):
                target.side_effect = [
                    InventoryVersionConflict(
                        "盘点任务已被其他设备更新，请刷新后重试",
                        current_version=9,
                    ),
                    success_value,
                ]
                stale = getattr(client, method)(
                    path, json={**body, "expected_version": 3}
                )
                self.assertEqual(stale.status_code, 409)
                self.assertEqual(stale.get_json()["current_version"], 9)
                retried = getattr(client, method)(
                    path, json={**body, "expected_version": 9}
                )
                self.assertEqual(retried.status_code, 200)
                self.assertEqual(retried.get_json()["version"], 10)
                self.assertIn(key, retried.get_json())
                target.side_effect = None

    def test_wrong_methods_are_rejected_for_every_inventory_route(self):
        client = self.login_account("admin", "admin-pass")
        cases = [
            ("post", "/api/inventory/tasks/active"),
            ("get", "/api/inventory/tasks"),
            ("post", "/api/inventory/tasks/history"),
            ("post", "/api/inventory/tasks/task-1"),
            ("get", "/api/inventory/tasks/task-1/items/A/claim"),
            ("get", "/api/inventory/tasks/task-1/items/A/heartbeat"),
            ("get", "/api/inventory/tasks/task-1/items/A/count"),
            ("get", "/api/inventory/tasks/task-1/items/A/serial/open"),
            ("get", "/api/inventory/tasks/task-1/items/A/serial/refresh"),
            ("get", "/api/inventory/tasks/task-1/items/A/serials"),
            ("post", "/api/inventory/tasks/task-1/items/A/serials/SN"),
            ("get", "/api/inventory/tasks/task-1/items/A/serial/finish"),
            ("delete", "/api/inventory/tasks/task-1/items/A/carton-preset"),
            ("get", "/api/inventory/tasks/task-1/items/A/cartons"),
            ("get", "/api/inventory/tasks/task-1/items/A/cartons/7/serials"),
            ("post", "/api/inventory/tasks/task-1/items/A/cartons/7/serials/SN"),
            ("post", "/api/inventory/tasks/task-1/items/A/cartons/7"),
            ("get", "/api/inventory/tasks/task-1/complete"),
            ("get", "/api/inventory/tasks/task-1/items/A/unlock"),
            ("post", "/api/inventory/tasks/task-1/export"),
            ("post", "/api/inventory/discrepancies"),
            ("get", "/api/inventory/discrepancies/1/notes"),
            ("get", "/api/inventory/discrepancies/1/archive"),
            ("get", "/api/inventory/discrepancies/1/restore"),
            ("post", "/api/inventory/discrepancies/export"),
        ]
        for method, path in cases:
            with self.subTest(method=method, path=path):
                self.assertEqual(getattr(client, method)(path).status_code, 405)

    def test_inventory_exception_statuses_and_worker_errors_are_sanitized(self):
        client = self.login_account("counter")
        cases = [
            (InventoryConflict("locked"), 409, "locked"),
            (InventoryNotFound("missing"), 404, "missing"),
            (InventoryPermissionDenied("admin only"), 403, "admin only"),
            (ValueError("bad quantity"), 400, "bad quantity"),
        ]
        for error, status, message in cases:
            with self.subTest(error=type(error).__name__):
                self.service.create_task.side_effect = error
                response = client.post("/api/inventory/tasks")
                self.assertEqual(response.status_code, status)
                self.assertEqual(response.get_json(), {"success": False, "error": message})

        secret = "成本价=999; password=secret; <html>private page</html>"
        self.service.create_task.side_effect = InventoryServiceError(secret)
        response = client.post("/api/inventory/tasks")
        self.assertEqual(response.status_code, 502)
        self.assertFalse(response.get_json()["success"])
        self.assertNotIn("999", response.get_data(as_text=True))
        self.assertNotIn("secret", response.get_data(as_text=True))
        self.assertNotIn("html", response.get_data(as_text=True).lower())

        self.store.get_active_task.return_value = {"task_id": "task-1"}
        self.store.get_task_snapshot.return_value = {
            "success": True,
            "task_id": "task-1",
            "phase": "sync_error",
            "version": 4,
            "gyj_status": secret,
            "items": [],
        }
        with mock.patch.object(app_module, "_ensure_inventory_sync"):
            response = client.get("/api/inventory/tasks/active")
        body = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("999", body)
        self.assertNotIn("secret", body)
        self.assertNotIn("html", body.lower())

    def test_carton_routes_use_authenticated_identity(self):
        client = self.login_account("counter")
        base = "/api/inventory/tasks/task-1/items/A%2FB"

        fetched = client.get(f"{base}/carton-preset")
        saved = client.post(
            f"{base}/carton-preset",
            json={"device_id": "device-a", "carton_quantity": 20},
        )
        created = client.post(
            f"{base}/cartons",
            json={
                "device_id": "device-a", "preset_quantity": 20,
                "start_serial": "SN/1", "serials": ["SN/1", "SN/2"],
            },
        )
        added = client.post(
            f"{base}/cartons/7/serials",
            json={"device_id": "device-b", "serial": "SN/3"},
        )
        removed = client.delete(
            f"{base}/cartons/7/serials/SN%2F2",
            json={"device_id": "device-c"},
        )
        deleted = client.delete(
            f"{base}/cartons/7", json={"device_id": "device-d"},
        )

        self.assertEqual(fetched.get_json()["preset"]["carton_quantity"], 20)
        for response in (saved, created, added, removed, deleted):
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()["version"], 4)
        self.store.get_carton_preset.assert_called_once_with(
            "admin", "task-1", "A/B"
        )
        self.service.save_carton_preset.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "counter", 20
        )
        self.service.create_serial_carton.assert_called_once_with(
            "admin", "task-1", "A/B", "device-a", "counter",
            20, "SN/1", ["SN/1", "SN/2"],
        )
        self.service.add_carton_serial.assert_called_once_with(
            "admin", "task-1", "A/B", "device-b", "counter", 7, "SN/3"
        )
        self.service.remove_carton_serial.assert_called_once_with(
            "admin", "task-1", "A/B", "device-c", "counter", 7, "SN/2"
        )
        self.service.delete_serial_carton.assert_called_once_with(
            "admin", "task-1", "A/B", "device-d", "counter", 7
        )

    def test_carton_routes_validate_methods_and_errors(self):
        client = self.login_account("counter")
        base = "/api/inventory/tasks/task-1/items/A/cartons"

        obsolete = client.post(base, json={
            "device_id": "device-a", "carton_code": "BOX-1",
            "preset_quantity": 20, "start_serial": "S001",
            "serials": ["S001"],
        })
        self.assertEqual(obsolete.status_code, 400)
        self.assertIn("请求字段", obsolete.get_json()["error"])
        self.service.create_serial_carton.assert_not_called()

        malformed = client.post(base, json=["not", "an", "object"])
        self.assertEqual(malformed.status_code, 400)
        self.assertEqual(client.get(base).status_code, 405)
        self.assertEqual(
            client.post(f"{base}/0/serials", json={
                "device_id": "device-a", "serial": "S001",
            }).status_code,
            400,
        )

        self.service.create_serial_carton.side_effect = InventoryConflict(
            "任务内序列号重复：S001"
        )
        conflict = client.post(base, json={
            "device_id": "device-a", "preset_quantity": 20,
            "start_serial": "S001", "serials": ["S001"],
        })
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(conflict.get_json()["error"], "任务内序列号重复：S001")

    def test_lock_owner_is_returned_only_for_real_owner_conflicts(self):
        client = self.login_account("counter")
        owned = InventoryConflict(
            "商品已被其他设备锁定", lock_owner="\x00乙\n操作员\t"
        )

        self.service.open_serial_item.side_effect = owned
        serial = client.post(
            "/api/inventory/tasks/task-1/items/A/serial/open",
            json={"device_id": "device-a", "expected_version": 3},
        )
        self.assertEqual(serial.status_code, 409)
        self.assertEqual(serial.get_json()["lock_owner"], "乙操作员")

        self.service.open_serial_item.side_effect = InventoryConflict(
            "当前阶段不允许此操作"
        )
        ordinary = client.post(
            "/api/inventory/tasks/task-1/items/A/serial/open",
            json={"device_id": "device-a", "expected_version": 3},
        )
        self.assertEqual(ordinary.status_code, 409)
        self.assertNotIn("lock_owner", ordinary.get_json())

    def test_real_service_maps_raised_worker_error_to_redacted_502(self):
        secret = "cost=999; password=secret; <html>private page</html>"

        class RaisingWorker:
            def load_inventory_catalog(self):
                raise RuntimeError(secret)

        store = InventoryStore(
            os.path.join(self.tempdir.name, "route-worker-error.sqlite3")
        )
        service = InventoryService(store, lambda _owner: RaisingWorker())
        with (
            mock.patch.object(app_module, "inventory_store", store),
            mock.patch.object(app_module, "inventory_service", service),
        ):
            response = self.login_account("counter").post("/api/inventory/tasks")

        body = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 502)
        self.assertEqual(
            response.get_json(),
            {
                "success": False,
                "error": "GYJ 库存读取失败，请检查登录状态后重试",
            },
        )
        self.assertNotIn("999", body)
        self.assertNotIn("secret", body)
        self.assertNotIn("html", body.lower())

    def test_real_service_background_sync_persists_error_and_active_retry_recovers(self):
        secret = "cost=999; password=secret; <html>private page</html>"
        retry_started = threading.Event()

        class RetryWorker:
            def __init__(self):
                self.calls = 0

            def read_inventory_stock_totals(self):
                self.calls += 1
                if self.calls == 1:
                    raise RuntimeError(secret)
                retry_started.set()
                return True, {"A1": "2"}

        store = InventoryStore(
            os.path.join(self.tempdir.name, "route-sync-retry.sqlite3")
        )
        task = store.create_task("admin", "counter", [{
            "barcode": "A1",
            "name": "Filter",
            "spec": "",
            "model": "",
            "category": "Parts",
            "unit": "piece",
            "has_serial": False,
            "initial_stock": "2",
        }])
        task_id = task["task_id"]
        store.advance_count_phase_if_ready("admin", task_id)
        store.claim_item(task_id, "A1", "device-a", "counter", "counting")
        store.set_open_book_quantity(
            "admin", task_id, "A1", "device-a", "counter", "2"
        )
        store.record_count(
            "admin",
            task_id,
            "A1",
            "device-a",
            "counter",
            completed_book_qty="2",
            completed_actual_qty="2",
            diff_qty="0",
            state="matched",
        )
        worker = RetryWorker()
        service = InventoryService(store, lambda _owner: worker)

        with (
            mock.patch.object(app_module, "inventory_store", store),
            mock.patch.object(app_module, "inventory_service", service),
        ):
            app_module._run_inventory_sync("admin", task_id)
            failed = store.get_task_snapshot("admin", task_id)
            self.assertEqual(failed["phase"], "counting")
            self.assertIsNone(failed["last_sync_at"])
            self.assertIsNotNone(failed["last_sync_attempt_at"])
            self.assertNotIn("999", failed["gyj_status"])
            self.assertNotIn("secret", failed["gyj_status"])
            self.assertNotIn("html", failed["gyj_status"].lower())

            response = self.login_account("counter").get(
                "/api/inventory/tasks/active"
            )
            self.assertEqual(response.status_code, 200)
            response_body = response.get_data(as_text=True)
            self.assertNotIn("999", response_body)
            self.assertNotIn("secret", response_body)
            self.assertNotIn("html", response_body.lower())
            with store.connect() as connection:
                connection.execute(
                    "UPDATE inventory_tasks SET last_sync_attempt_at = ? "
                    "WHERE task_id = ?",
                    ((datetime.now() - timedelta(seconds=60)).isoformat(), task_id),
                )
            response = self.login_account("counter").get(
                "/api/inventory/tasks/active"
            )
            self.assertEqual(response.status_code, 200)
            self.assertTrue(retry_started.wait(1))
            with app_module.inventory_sync_lock:
                retry_thread = app_module.inventory_sync_threads.get(task_id)
            if retry_thread is not None:
                retry_thread.join(1)

        recovered = store.get_task_snapshot("admin", task_id)
        self.assertEqual(worker.calls, 2)
        self.assertEqual(recovered["phase"], "counting")
        self.assertEqual(recovered["gyj_status"], "synced")
        self.assertIsNotNone(recovered["last_sync_at"])

    def test_invalid_version_quantity_and_missing_device_are_400(self):
        self.store.get_active_task.return_value = {"task_id": "task-1"}
        client = self.login_account("counter")
        self.assertEqual(
            client.get("/api/inventory/tasks/active?version=not-an-int").status_code,
            400,
        )
        self.assertEqual(
            client.post(
                "/api/inventory/tasks/task-1/items/A/count-entries",
                json={"device_id": "device-a", "quantity": "not-a-number"},
            ).status_code,
            400,
        )
        self.assertEqual(
            client.post(
                "/api/inventory/tasks/task-1/items/A/count-entries",
                json={"quantity": "1"},
            ).status_code,
            400,
        )
        self.assertEqual(
            client.post(
                "/api/inventory/tasks/task-1/items/A/count-entries", json=[]
            ).status_code,
            400,
        )
        self.assertEqual(
            client.get(
                "/api/inventory/discrepancies?state=pending"
            ).status_code,
            400,
        )

    def test_non_admin_cannot_unlock_archive_or_restore_even_with_spoofed_json(self):
        client = self.login_account("counter")
        paths = [
            "/api/inventory/tasks/task-1/items/A/unlock",
            "/api/inventory/discrepancies/1/archive",
            "/api/inventory/discrepancies/1/restore",
        ]
        for path in paths:
            with self.subTest(path=path):
                response = client.post(path, json={"is_admin": True, "actor": "admin"})
                self.assertEqual(response.status_code, 403)
                self.assertFalse(response.get_json()["success"])
        self.store.admin_unlock.assert_not_called()
        self.store.archive_discrepancy.assert_not_called()
        self.store.restore_discrepancy.assert_not_called()

    def test_only_admin_can_reopen_inventory_task(self):
        ordinary = self.login_account("counter")
        denied = ordinary.post(
            "/api/inventory/tasks/task-1/reopen", json={"is_admin": True}
        )
        self.assertEqual(denied.status_code, 403)
        self.service.reopen_task.assert_not_called()

        admin = self.login_account("admin", "admin-pass")
        reopened = admin.post("/api/inventory/tasks/task-1/reopen", json={})
        self.assertEqual(reopened.status_code, 200)
        self.assertEqual(reopened.get_json(), {
            "success": True,
            "task": {"task_id": "task-1", "phase": "counting", "version": 9},
            "version": 9,
        })
        self.service.reopen_task.assert_called_once_with(
            "admin", "task-1", "admin"
        )

    def test_only_admin_can_delete_completed_inventory_task(self):
        ordinary = self.login_account("counter")
        denied = ordinary.delete("/api/inventory/tasks/task-1")
        self.assertEqual(denied.status_code, 403)
        self.store.delete_completed_task.assert_not_called()

        self.store.delete_completed_task.return_value = {"task_id": "task-1"}
        admin = self.login_account("admin", "admin-pass")
        deleted = admin.delete("/api/inventory/tasks/task-1")
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(deleted.get_json(), {
            "success": True, "deleted": {"task_id": "task-1"},
        })
        self.store.delete_completed_task.assert_called_once_with(
            "admin", "task-1"
        )

    def test_audit_route_is_owner_scoped_and_filters_barcode(self):
        response = self.login_account("counter").get(
            "/api/inventory/tasks/task-1/audit?barcode=A%2FB"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["events"][0]["event_label"], "修改分次数量")
        self.store.list_audit_events.assert_called_once_with(
            "admin", "task-1", barcode="A/B"
        )

    def test_history_and_discrepancy_route_json_contract_matches_store_rows(self):
        history_row = {
            "task_id": "task-history-1", "owner": "admin",
            "created_by": "盘点员", "phase": "completed", "version": 8,
            "started_at": "2026-09-01T08:00:00",
            "completed_at": "2026-09-01T09:00:00", "last_sync_at": None,
            "gyj_status": "synced", "sync_resume_phase": None,
            "completed": True,
            "product_total": 4, "quantity_difference_count": 2,
            "serial_difference_count": 3, "participant_count": 2,
            "counted_product_count": 3, "uncounted_product_count": 1,
        }
        discrepancy_row = {
            "id": 9, "task_id": "task-history-1", "barcode": "A/B",
            "serial": "SN/1", "kind": "other_product_serial",
            "book_quantity": None, "counted_quantity": None,
            "difference": None, "state": "open", "archived_by": None,
            "archived_at": None, "created_at": "2026-09-01T09:00:00",
            "name": "滤芯", "spec": "S", "model": "M", "category": "耗材",
            "unit": "支", "has_serial": True,
            "completed_at": "2026-09-01T09:00:00", "scan_actor": "盘点员",
            "scan_device": "device-a", "scanned_at": "2026-09-01T08:30:00",
            "lookup_barcode": "OTHER", "lookup_name": "其他商品",
            "warehouse": "南昌仓", "shipped": False,
            "notes": [{
                "id": 1, "task_id": "task-history-1", "barcode": "A/B",
                "serial": "SN/1", "note": "已复核", "actor": "盘点员",
                "created_at": "2026-09-01T09:10:00",
            }],
            "note": "已复核",
        }
        self.store.list_task_history.return_value = {
            "tasks": [history_row], "total": 21, "limit": 10, "offset": 10,
        }
        self.store.list_discrepancies.return_value = [discrepancy_row]
        client = self.login_account("counter")

        history = client.get(
            "/api/inventory/tasks/history?limit=10&offset=10"
        ).get_json()
        discrepancies = client.get(
            "/api/inventory/discrepancies?state=open&query=SN%2F1"
        ).get_json()

        self.assertEqual(history, {
            "success": True,
            "tasks": [{**history_row, "task_number": "PD20260901-090000"}],
            "pagination": {
                "limit": 10, "offset": 10, "total": 21, "has_more": True,
            },
        })
        self.store.list_task_history.assert_called_with(
            "admin", limit=10, offset=10
        )
        self.assertEqual(
            discrepancies,
            {"success": True, "discrepancies": [{
                **discrepancy_row, "task_number": "PD20260901-090000",
            }]},
        )
        self.store.list_discrepancies.assert_called_with(
            "admin", "open", query="SN/1"
        )

    def test_history_route_rejects_unbounded_or_invalid_pagination(self):
        client = self.login_account("counter")
        for query in ("limit=0", "limit=51", "limit=nope", "offset=-1", "offset=nope"):
            with self.subTest(query=query):
                response = client.get(f"/api/inventory/tasks/history?{query}")
                self.assertEqual(response.status_code, 400)
                self.assertFalse(response.get_json()["success"])
        self.store.list_task_history.assert_not_called()

    def test_history_task_number_and_detail_route_use_completion_time(self):
        task = {
            "task_id": "task-history-1", "owner": "admin",
            "created_by": "盘点员", "phase": "completed", "version": 8,
            "started_at": "2026-09-07T08:00:00",
            "completed_at": "2026-09-07T08:14:26.123456",
            "last_sync_at": None, "gyj_status": "synced",
            "sync_resume_phase": None,
        }
        self.store.list_task_history.return_value = {
            "tasks": [{**task, "completed": True, "product_total": 1,
                       "counted_product_count": 1,
                       "uncounted_product_count": 0,
                       "quantity_difference_count": 1,
                       "serial_difference_count": 1,
                       "participant_count": 1}],
            "total": 1, "limit": 20, "offset": 0,
        }
        self.store.get_task_history_detail.return_value = {
            "task": task, "scope": "counted", "participants": [],
            "items": [{
                "barcode": "B2", "name": "序列号商品", "has_serial": True,
                "completed_book_qty": "2", "completed_actual_qty": "1",
                "diff_qty": "-1", "serial_discrepancies": [{
                    "id": 9, "serial": "SN-1", "kind": "system_only_serial",
                    "state": "archived", "archived_by": "管理员",
                    "archived_at": "2026-09-07T09:00:00",
                }],
            }],
        }
        client = self.login_account("counter")

        history = client.get("/api/inventory/tasks/history").get_json()
        detail_response = client.get(
            "/api/inventory/tasks/task-history-1/history-detail?scope=counted"
        )

        self.assertEqual(history["tasks"][0]["task_number"], "PD20260907-081426")
        self.assertEqual(detail_response.status_code, 200)
        detail = detail_response.get_json()
        self.assertEqual(detail["task"]["task_number"], "PD20260907-081426")
        self.assertEqual(detail["scope"], "counted")
        self.assertEqual(detail["items"][0]["serial_discrepancies"][0]["state"], "archived")
        self.store.get_task_history_detail.assert_called_once_with(
            "admin", "task-history-1", scope="counted"
        )

    def test_history_detail_route_rejects_unknown_metric_scope(self):
        response = self.login_account("counter").get(
            "/api/inventory/tasks/task-history-1/history-detail?scope=prices"
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.get_json()["success"])
        self.store.get_task_history_detail.assert_not_called()

    def test_export_routes_are_owner_scoped_filtered_and_price_free(self):
        self.store.get_task_snapshot.return_value = {
            "success": True,
            "task_id": "task-unsafe\r\nname",
            "phase": "completed",
            "items": [{
                "barcode": "A1", "name": "Filter", "diff_qty": "1",
                "completed_book_qty": "1", "completed_actual_qty": "2",
                "purchase_price": "SECRET-PRICE",
            }],
        }
        self.store.list_discrepancies.side_effect = [
            [{"task_id": "task-unsafe\r\nname", "barcode": "A1", "serial": None, "state": "open", "cost": "SECRET-PRICE"}],
            [],
            [{"task_id": "task-unsafe\r\nname", "barcode": "A1", "serial": "SN1", "state": "open", "price": "SECRET-PRICE"}],
        ]
        client = self.login_account("counter")
        task_response = client.get("/api/inventory/tasks/task-safe/export")
        self.assertEqual(task_response.status_code, 200)
        self.assertNotIn("\r", task_response.headers["Content-Disposition"])
        self.assertNotIn("\n", task_response.headers["Content-Disposition"])
        workbook = load_workbook(BytesIO(task_response.data))
        self.assertNotIn("SECRET-PRICE", repr([
            cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
        ]))
        self.store.get_task_snapshot.assert_called_with(
            "admin", "task-safe"
        )

        discrepancy_response = client.get(
            "/api/inventory/discrepancies/export?state=open&query=SN1"
        )
        self.assertEqual(discrepancy_response.status_code, 200)
        discrepancy_workbook = load_workbook(BytesIO(discrepancy_response.data))
        self.assertNotIn("SECRET-PRICE", repr([
            cell.value
            for sheet in discrepancy_workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ]))
        self.store.list_discrepancies.assert_called_with(
            "admin", "open", query="SN1"
        )

    def test_ensure_inventory_sync_deduplicates_daemons_and_reopens_stale_tasks(self):
        stale = (datetime.now() - timedelta(seconds=61)).isoformat(timespec="seconds")
        self.store.get_task_snapshot.return_value = {
            "task_id": "task-1", "phase": "counting", "last_sync_at": stale
        }
        started = threading.Event()
        release = threading.Event()
        calls = []

        def sync(owner, task_id):
            calls.append((owner, task_id))
            started.set()
            release.wait(2)

        self.service.sync_completed_items.side_effect = sync
        self.assertTrue(app_module._ensure_inventory_sync("counter-id", "task-1"))
        self.assertTrue(started.wait(1))
        thread = app_module.inventory_sync_threads["task-1"]
        self.assertTrue(thread.daemon)
        self.assertFalse(app_module._ensure_inventory_sync("counter-id", "task-1"))
        self.assertEqual(calls, [("counter-id", "task-1")])

        release.set()
        thread.join(1)
        self.assertNotIn("task-1", app_module.inventory_sync_threads)

        started.clear()
        release.clear()
        self.assertTrue(app_module._ensure_inventory_sync("counter-id", "task-1"))
        self.assertTrue(started.wait(1))
        self.assertEqual(len(calls), 2)
        release.set()

    def test_recent_sync_attempt_does_not_start_a_daemon_until_sixty_seconds(self):
        self.store.get_task_snapshot.return_value = {
            "task_id": "task-1",
            "phase": "counting",
            "last_sync_at": None,
            "last_sync_attempt_at": datetime.now().isoformat(timespec="seconds"),
        }
        self.assertFalse(app_module._ensure_inventory_sync("counter-id", "task-1"))
        self.service.sync_completed_items.assert_not_called()

        done = threading.Event()
        self.store.get_task_snapshot.return_value["last_sync_attempt_at"] = (
            datetime.now() - timedelta(seconds=60)
        ).isoformat(timespec="seconds")
        self.service.sync_completed_items.side_effect = lambda *_args: done.set()
        self.assertTrue(app_module._ensure_inventory_sync("counter-id", "task-1"))
        self.assertTrue(done.wait(1))

    def test_serial_prefetch_deduplicates_running_daemons_and_cleans_registry(self):
        started = threading.Event()
        release = threading.Event()

        def refresh(*_args, **_kwargs):
            started.set()
            release.wait(2)

        self.service.refresh_serial_item.side_effect = refresh
        self.assertTrue(app_module._ensure_inventory_serial_prefetch(
            "counter-id", "task-1", "B2", "device-a"
        ))
        self.assertTrue(started.wait(1))
        key = ("counter-id", "task-1", "B2")
        thread = app_module.inventory_serial_prefetch_threads[key]
        self.assertTrue(thread.daemon)
        items = [{"barcode": "B2", "has_serial": True}]
        app_module._inventory_attach_serial_prefetch_status(
            "counter-id", "task-1", items
        )
        self.assertTrue(items[0]["serial_syncing"])
        self.assertFalse(app_module._ensure_inventory_serial_prefetch(
            "counter-id", "task-1", "B2", "device-b"
        ))

        release.set()
        thread.join(1)
        self.assertNotIn(key, app_module.inventory_serial_prefetch_threads)
        app_module._inventory_attach_serial_prefetch_status(
            "counter-id", "task-1", items
        )
        self.assertFalse(items[0]["serial_syncing"])
        self.service.refresh_serial_item.assert_called_once_with(
            "counter-id", "task-1", "B2", "device-a", "system", force=False
        )

    def test_serial_prefetch_failure_is_nonblocking_and_cleans_registry(self):
        done = threading.Event()

        def fail(*_args, **_kwargs):
            done.set()
            raise InventoryServiceError("GYJ unavailable")

        self.service.refresh_serial_item.side_effect = fail
        self.assertTrue(app_module._ensure_inventory_serial_prefetch(
            "counter-id", "task-1", "B2", "device-a"
        ))
        self.assertTrue(done.wait(1))
        for thread in list(app_module.inventory_serial_prefetch_threads.values()):
            thread.join(1)
        self.assertNotIn(
            ("counter-id", "task-1", "B2"),
            app_module.inventory_serial_prefetch_threads,
        )

    def test_sync_thread_start_failure_cleans_registry_without_blocking_caller(self):
        self.store.get_task_snapshot.return_value = {
            "task_id": "task-1", "phase": "counting", "last_sync_at": None
        }
        thread = mock.Mock()
        thread.start.side_effect = RuntimeError("thread unavailable")
        with mock.patch.object(app_module.threading, "Thread", return_value=thread):
            self.assertFalse(
                app_module._ensure_inventory_sync("counter-id", "task-1")
            )
        self.assertNotIn("task-1", app_module.inventory_sync_threads)


if __name__ == "__main__":
    unittest.main()
