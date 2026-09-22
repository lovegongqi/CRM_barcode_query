import unittest

from openpyxl import load_workbook

from inventory_export import build_inventory_workbook, build_discrepancy_workbook


class InventoryExportTests(unittest.TestCase):
    def setUp(self):
        self.task = {"task_id": "T-1", "phase": "completed"}
        self.items = [{
            "barcode": "00123",
            "name": "滤芯",
            "spec": "标准",
            "model": "M1",
            "category": "配件",
            "unit": "个",
            "has_serial": True,
            "completed_book_qty": "10.00",
            "completed_actual_qty": "9.5",
            "diff_qty": "-0.5",
            "state": "variance",
            "note": "=备注内容",
            "成本价": "999",
            "库存金额": "9999",
        }]
        self.discrepancies = [{
            "task_id": "T-1", "barcode": "00123", "name": "滤芯",
            "serial": "SN-001", "kind": "physical_only_serial", "state": "open",
            "scan_actor": "张三", "scan_device": "设备A", "scanned_at": "2026-09-01T10:00:00",
            "note": "已找到，放错仓位", "采购价": "888",
        }]

    def test_inventory_workbook_has_exact_sheets_headers_and_no_price_leak(self):
        stream = build_inventory_workbook(self.task, self.items, self.discrepancies)
        self.assertEqual(stream.tell(), 0)
        workbook = load_workbook(stream, data_only=True)
        self.assertEqual(workbook.sheetnames, ["商品差异汇总", "序列号差异明细"])
        self.assertEqual([cell.value for cell in workbook.worksheets[0][1]], [
            "商品条码", "商品名称", "规格", "型号", "类别", "单位", "序列号管理",
            "完成时账面数量", "实盘数量", "盘盈/盘亏", "处理状态", "商品备注",
        ])
        self.assertEqual([cell.value for cell in workbook.worksheets[1][1]], [
            "商品条码", "商品名称", "序列号", "差异分类", "盘点任务", "扫描账号",
            "扫描设备", "扫描时间", "处理状态", "序列号备注",
        ])
        values = "|".join(str(cell.value or "") for sheet in workbook for row in sheet.iter_rows() for cell in row)
        for forbidden in ("成本价", "采购价", "零售价", "销售价", "库存金额", "入库单价", "999", "888"):
            self.assertNotIn(forbidden, values)
        self.assertIn("已找到，放错仓位", values)
        self.assertNotEqual(workbook.worksheets[0]["L2"].data_type, "f")
        self.assertEqual(workbook.worksheets[0]["A2"].value, "00123")
        self.assertEqual(workbook.worksheets[1]["C2"].value, "SN-001")

    def test_empty_and_unchanged_items_keep_headers_only(self):
        stream = build_inventory_workbook({}, [{"barcode": "A", "diff_qty": "0"}], [])
        workbook = load_workbook(stream)
        self.assertEqual(workbook.worksheets[0].max_row, 1)
        self.assertEqual(workbook.worksheets[1].max_row, 1)

    def test_discrepancy_workbook_keeps_state_and_is_rewound(self):
        stream = build_discrepancy_workbook(self.discrepancies, "已归档")
        self.assertEqual(stream.tell(), 0)
        workbook = load_workbook(stream, data_only=True)
        self.assertEqual(workbook.sheetnames, ["差异明细"])
        self.assertEqual(workbook["差异明细"]["I2"].value, "已归档")
        self.assertEqual(workbook["差异明细"]["C2"].number_format, "@")

    def test_unverified_serial_is_labeled_in_both_sheets_and_uncounted_is_omitted(self):
        items = [
            {
                "barcode": "SERIAL", "name": "序列号商品", "has_serial": True,
                "completed_book_qty": "5", "completed_actual_qty": "4",
                "diff_qty": "-1", "state": "serial_pending",
            },
            {
                "barcode": "UNCOUNTED", "name": "未盘商品", "has_serial": False,
                "completed_book_qty": None, "completed_actual_qty": None,
                "diff_qty": None, "state": "pending",
            },
        ]
        discrepancies = [
            {
                "task_id": "T-1", "barcode": "SERIAL", "name": "序列号商品",
                "serial": None, "kind": "serial_unverified", "state": "open",
                "book_quantity": "5", "counted_quantity": "4", "difference": "-1",
            },
            {
                "task_id": "T-1", "barcode": "SERIAL", "name": "序列号商品",
                "serial": None, "kind": "product_quantity", "state": "open",
                "book_quantity": "5", "counted_quantity": "4", "difference": "-1",
            },
        ]

        workbook = load_workbook(
            build_inventory_workbook(self.task, items, discrepancies),
            data_only=True,
        )
        summary_values = [cell.value for cell in workbook["商品差异汇总"][2]]
        serial_values = [cell.value for cell in workbook["序列号差异明细"][2]]
        self.assertIn("序列号未核对", summary_values)
        self.assertIn("序列号未核对", serial_values)
        all_values = "|".join(
            str(cell.value or "")
            for sheet in workbook for row in sheet.iter_rows() for cell in row
        )
        self.assertNotIn("UNCOUNTED", all_values)
        for forbidden in ("价格", "金额", "成本价", "采购价"):
            self.assertNotIn(forbidden, all_values)


if __name__ == "__main__":
    unittest.main()
